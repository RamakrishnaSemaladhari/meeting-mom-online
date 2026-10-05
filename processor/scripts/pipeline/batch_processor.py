"""Automatic 10-minute batch processing for long online meetings."""
import json
import os
import subprocess
from pathlib import Path

from .errors import MomError
from . import ai_understanding as ai_mod
from .docx_builder import build_mom_doc
from .validation import validate

BATCH_SECONDS = int(env("BATCH_SECONDS", "600"))

def _ffprobe_duration(path):
    try:
        proc = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
            capture_output=True, text=True, check=True, timeout=120)
        return float(proc.stdout.strip() or 0)
    except Exception as exc:
        raise MomError("MOM-004", f"Could not determine audio duration: {exc}")

def _shift_segments(segments, offset):
    return [{"start": round(s["start"] + offset, 2),
             "end": round(s["end"] + offset, 2), "text": s["text"]} for s in segments]

def _combined_bilingual(original, translation):
    lines = []
    for seg in original:
        lines.append(f"[{fmt_ts(seg['start'])} - {fmt_ts(seg['end'])}] ORIGINAL: {seg['text']}")
        lo, hi = seg["start"] - 1.0, seg["end"] + 1.0
        picked = [t for t in translation
                  if lo <= (t["start"] + t["end"]) / 2 <= hi]
        if picked:
            lines.append(f"[{fmt_ts(seg['start'])} - {fmt_ts(seg['end'])}] ENGLISH: {' '.join(t['text'] for t in picked)}")
        else:
            lines.append(f"[{fmt_ts(seg['start'])} - {fmt_ts(seg['end'])}] ENGLISH: [No aligned translation]")
        lines.append("")
    return "\n".join(lines).strip() + "\n"

def _build_bilingual_docx(path, title, original, translation):
    from docx import Document
    doc = Document()
    doc.add_heading(title or "Original + English Translation", 0)
    doc.add_paragraph("Timestamp-aligned original-language transcript and English translation. Original wording is preserved.")
    for seg in original:
        doc.add_paragraph(f"[{fmt_ts(seg['start'])} - {fmt_ts(seg['end'])}] ORIGINAL: {seg['text']}")
        lo, hi = seg["start"] - 1.0, seg["end"] + 1.0
        picked = [t for t in translation
                  if lo <= (t["start"] + t["end"]) / 2 <= hi]
        if picked:
            doc.add_paragraph(f"[{fmt_ts(seg['start'])} - {fmt_ts(seg['end'])}] ENGLISH: {' '.join(t['text'] for t in picked)}")
        else:
            doc.add_paragraph(f"[{fmt_ts(seg['start'])} - {fmt_ts(seg['end'])}] ENGLISH: [No aligned translation]")
    doc.save(path)

def _shift_whisper_data(data, offset):
    out = dict(data)
    rows = []
    for row in data.get("transcription") or data.get("segments") or []:
        item = dict(row)
        if item.get("offsets") and "from" in item["offsets"]:
            offsets = dict(item["offsets"])
            offsets["from"] = int(offsets.get("from", 0)) + int(offset * 1000)
            if "to" in offsets:
                offsets["to"] = int(offsets.get("to", 0)) + int(offset * 1000)
            item["offsets"] = offsets
        else:
            item["start"] = round(float(item.get("start", 0)) + offset, 2)
            item["end"] = round(float(item.get("end", item.get("start", 0))) + offset, 2)
        rows.append(item)
    if "transcription" in out:
        out["transcription"] = rows
    elif "segments" in out:
        out["segments"] = rows
    return out

def _read_previous_context(ctx):
    try:
        meta = ctx.store.get_meta(ctx.meeting_folder)
        parent = (meta.get("parents") or [None])[0]
        if not parent:
            return None
        siblings = [x for x in ctx.store.list_children(parent, folders=True)
                    if x["id"] != ctx.meeting_folder]
        siblings.sort(key=lambda x: x.get("createdTime", ""), reverse=True)
        for sibling in siblings[:10]:
            children = ctx.store.list_children(sibling["id"], folders=True)
            ai_folder = next((x for x in children if x["name"] == "AI"), None)
            if not ai_folder:
                continue
            evidence = ctx.store.find_one(ai_folder["id"], "AI Evidence.json")
            if not evidence:
                continue
            data = json.loads(ctx.store.read_bytes(evidence["id"]).decode("utf-8"))
            return {
                "meeting_name": sibling.get("name", "Previous meeting"),
                "executive_summary": data.get("executive_summary", data.get("summary", "")),
                "decisions": data.get("decisions", []),
                "action_items": data.get("action_items", []),
                "open_questions": data.get("open_questions", [])
            }
    except Exception as exc:
        log(f"Previous-meeting context ignored: {exc}")
    return None

def _aggregate_batches(metadata, batch_results, text_fn, previous=None, model=""):
    sections = []
    for b in batch_results:
        sections.append(ai_mod.normalize_chunk({
            "section_summary": b.get("complete_conversation_summary", ""),
            "discussion_points": b.get("discussion_points", []),
            "decisions": b.get("decisions", []),
            "action_items": b.get("action_items", []),
            "commitments": b.get("commitments", []),
            "open_questions": b.get("open_questions", []),
            "next_meeting": b.get("next_meeting", []),
            "review_flags": b.get("review_flags", [])
        }))
    merged = ai_mod.merge_sections(sections)
    summaries = [s["section_summary"] for s in sections if s["section_summary"]]
    context_note = ""
    if previous:
        context_note = (
            "\nOPTIONAL PREVIOUS MEETING CONTEXT. Use it only if the current meeting explicitly "
            "continues or refers to it; otherwise ignore it:\n" +
            json.dumps(previous, ensure_ascii=False)[:12000]
        )
    complete = summaries[0] if len(summaries) == 1 else text_fn(
        ai_mod.complete_summary_prompt(metadata, summaries) + context_note
    )
    executive = text_fn(
        ai_mod.executive_summary_prompt(metadata, complete, merged) + context_note
    ) if complete else ""
    result = {"executive_summary": executive, "complete_conversation_summary": complete}
    result.update(merged)
    result = {k: result[k] for k in ai_mod.OUTPUT_KEYS}
    result["schema_version"] = ai_mod.SCHEMA_VERSION
    result["generated_with"] = {
        "ai_model": model, "batches": len(batch_results), "batch_seconds": BATCH_SECONDS
    }
    result["previous_meeting_context"] = {
        "used_as_reference": bool(previous),
        "meeting_name": previous.get("meeting_name", "") if previous else ""
    }
    return result

def stage_batch(ctx):
    st = ctx.status
    metadata = load_metadata(ctx)
    audio, warnings = select_audio(ctx.store, ctx.audio_folder, env("AUDIO_FILE_ID"), True)
    st.update("DOWNLOADING", "Downloading the full meeting audio", 5,
              batch_index=0, batch_total=0)
    ctx.store.download(audio["id"], ctx.p["source"])
    duration = _ffprobe_duration(ctx.p["source"])
    if duration <= 0:
        raise MomError("MOM-004", "The uploaded audio has no measurable duration.")
    total = int((duration + BATCH_SECONDS - 1) // BATCH_SECONDS)
    if total > 60:
        raise MomError("MOM-004", "Meeting exceeds the 60-batch safety limit.")
    batches_folder = ctx.store.ensure_folder(ctx.meeting_folder, "BATCHES")
    ctx.store.upsert_text(ctx.meeting_folder, "BATCH_MANIFEST.json", dump({
        "meeting_id": env("MEETING_ID"), "duration_seconds": duration,
        "batch_seconds": BATCH_SECONDS, "batch_total": total, "status": "PROCESSING"
    }), JSON_MIME)

    chunks_dir = WORK / "batches"
    chunks_dir.mkdir(parents=True, exist_ok=True)
    run_cmd([
        "ffmpeg", "-y", "-i", str(ctx.p["source"]), "-map", "0:a:0",
        "-f", "segment", "-segment_time", str(BATCH_SECONDS),
        "-reset_timestamps", "1", "-ar", "16000", "-ac", "1",
        str(chunks_dir / "batch_%03d.wav")
    ], "MOM-004")
    files = sorted(chunks_dir.glob("batch_*.wav"))
    if not files:
        raise MomError("MOM-004", "No 10-minute audio batches were created.")

    model = env("AI_MODEL", "qwen3:4b-instruct-2507-q4_K_M")
    json_fn, text_fn = ai_mod.make_ollama_fns(model)
    batch_results, all_original, all_translation = [], [], []
    previous = _read_previous_context(ctx)

    for idx, chunk_file in enumerate(files, 1):
        offset = (idx - 1) * BATCH_SECONDS
        end_time = min(offset + BATCH_SECONDS, duration)
        batch_name = f"BATCH_{idx:03d}_{fmt_ts(offset).replace(':','-')}_{fmt_ts(end_time).replace(':','-')}"
        batch_folder = ctx.store.ensure_folder(batches_folder, batch_name)

        st.update("TRANSCRIBING", f"Batch {idx}/{len(files)} — Whisper transcription",
                  10 + int(42 * (idx - 1) / len(files)),
                  batch_index=idx, batch_total=len(files))
        seconds = min(BATCH_SECONDS, max(1, duration - offset))
        data, local_segments = transcribe_with_gate(
            chunk_file, f"{chunks_dir}/transcript_{idx:03d}", seconds,
            language_hint(metadata.get("language_hint", "")), st
        )
        original = _shift_segments(local_segments, offset)
        data = _shift_whisper_data(data, offset)
        data["_batch"] = {"index": idx, "start_seconds": offset, "end_seconds": end_time}
        all_original.extend(original)
        ctx.store.upsert_text(batch_folder, "Original Transcript.txt", timestamped_text(original))
        ctx.store.upsert_text(batch_folder, "Transcript.json", dump(data), JSON_MIME)

        st.update("TRANSLATING", f"Batch {idx}/{len(files)} — English translation",
                  52 + int(15 * (idx - 1) / len(files)),
                  batch_index=idx, batch_total=len(files))
        lang = detected_language(data)
        if lang == "en":
            translation = [dict(s) for s in original]
            tdata = {"transcription": [
                {"offsets": {"from": int(s["start"]*1000), "to": int(s["end"]*1000)}, "text": s["text"]}
                for s in original
            ], "result": {"language": "en"},
               "_quality": {"passed": True, "skipped": "source language is English"}}
        else:
            tdata_local = whisper_json(
                chunk_file, f"{chunks_dir}/translation_{idx:03d}",
                translate=True, model=WHISPER_MODEL, language=lang or "auto"
            )
            translation = _shift_segments(ai_mod.segments_from_whisper(tdata_local), offset)
            tdata = _shift_whisper_data(tdata_local, offset)
            tdata["_batch"] = data["_batch"]
        all_translation.extend(translation)
        ctx.store.upsert_text(
            batch_folder, "English Translation.txt",
            timestamped_text(translation) if translation else "No English translation produced.\n"
        )
        ctx.store.upsert_text(batch_folder, "Translation.json", dump(tdata), JSON_MIME)
        ctx.store.upsert_text(
            batch_folder, "Original + English Translation.txt",
            _combined_bilingual(original, translation)
        )

        st.update("ANALYZING", f"Batch {idx}/{len(files)} — AI evidence extraction",
                  68 + int(12 * (idx - 1) / len(files)),
                  batch_index=idx, batch_total=len(files))
        ai = ai_mod.run_understanding(
            metadata, original, translation, json_fn, text_fn,
            max_chars=int(env("AI_CHUNK_CHARS", "8000")), model=model
        )
        ai["batch"] = {"index": idx, "start": fmt_ts(offset), "end": fmt_ts(end_time)}
        ctx.store.upsert_text(batch_folder, "AI Evidence.json", dump(ai), JSON_MIME)
        ctx.store.upsert_text(batch_folder, "Complete Conversation Summary.txt",
                              ai.get("complete_conversation_summary", ""))
        ctx.store.upsert_text(batch_folder, "Executive Summary.txt",
                              ai.get("executive_summary", ""))
        batch_results.append(ai)

    st.update("SUMMARIZING", "Combining all batches into one complete meeting understanding",
              84, batch_index=len(files), batch_total=len(files))
    final_ai = _aggregate_batches(metadata, batch_results, text_fn, previous=previous, model=model)
    final_ai["batch_manifest"] = {
        "total": len(files), "batch_seconds": BATCH_SECONDS, "duration_seconds": duration
    }
    ctx.p["ai"].write_text(dump(final_ai), encoding="utf-8")
    publish_ai_outputs(ctx, final_ai, validated=False)

    st.update("GENERATING_MOM", "Validating evidence across the complete meeting", 90,
              batch_index=len(files), batch_total=len(files))
    try:
        validated, report = validate(final_ai, all_original, all_translation)
    except Exception as exc:
        raise MomError("MOM-011", f"Evidence validation failed: {exc}")

    st.update("GENERATING_MOM", "Building the single meeting MoM", 94,
              batch_index=len(files), batch_total=len(files))
    try:
        build_mom_doc(metadata, validated).save(ctx.p["mom"])
    except Exception as exc:
        raise MomError("MOM-012", f"DOCX generation failed: {exc}")

    st.update("UPLOADING", "Uploading complete transcript, translation, AI evidence and single MoM",
              97, batch_index=len(files), batch_total=len(files))
    publish_ai_outputs(ctx, validated, validated=True)
    title = metadata.get("title") or "Meeting"
    docx_name = f"{title} - AI_MOM.docx"
    ctx.store.upsert_file(ctx.folders["MOM"], docx_name, ctx.p["mom"], DOCX_MIME)

    ctx.store.upsert_text(
        ctx.folders["TRANSCRIPT"], "Original Transcript - Complete Meeting.txt",
        timestamped_text(all_original)
    )
    ctx.store.upsert_text(
        ctx.folders["TRANSLATION"], "English Translation - Complete Meeting.txt",
        timestamped_text(all_translation)
    )
    ctx.store.upsert_text(
        ctx.folders["TRANSLATION"], "Original + English Translation - Complete Meeting.txt",
        _combined_bilingual(all_original, all_translation)
    )
    bilingual_path = WORK / "Original + English Translation - Complete Meeting.docx"
    _build_bilingual_docx(
        bilingual_path, metadata.get("title") or "Original + English Translation",
        all_original, all_translation
    )
    ctx.store.upsert_file(
        ctx.folders["TRANSLATION"],
        "Original + English Translation - Complete Meeting.docx",
        bilingual_path, DOCX_MIME
    )
    ctx.store.upsert_text(
        ctx.folders["TRANSLATION"], "Translation.json",
        dump({"transcription": [
            {"offsets": {"from": int(s["start"]*1000), "to": int(s["end"]*1000)},
             "text": s["text"]} for s in all_translation
        ], "result": {"language": "en"}}), JSON_MIME
    )
    ctx.store.upsert_text(
        ctx.folders["TRANSCRIPT"], "Transcript.json",
        dump({"transcription": [
            {"offsets": {"from": int(s["start"]*1000), "to": int(s["end"]*1000)},
             "text": s["text"]} for s in all_original
        ]}), JSON_MIME
    )

    ctx.p["transcript_json"].write_text(dump({"transcription": [
        {"offsets": {"from": int(s["start"]*1000), "to": int(s["end"]*1000)}, "text": s["text"]}
        for s in all_original
    ]}), encoding="utf-8")
    ctx.p["translation_json"].write_text(dump({"transcription": [
        {"offsets": {"from": int(s["start"]*1000), "to": int(s["end"]*1000)}, "text": s["text"]}
        for s in all_translation
    ], "result": {"language": "en"}}), encoding="utf-8")

    if ctx.p["warnings"].exists():
        warnings.extend(json.loads(ctx.p["warnings"].read_text(encoding="utf-8")))
    ctx.store.upsert_text(ctx.meeting_folder, "PROCESSING_COMPLETE.json", dump({
        "meeting_id": env("MEETING_ID"), "status": "complete", "completed_at": now(),
        "batch_total": len(files), "batch_seconds": BATCH_SECONDS,
        "duration_seconds": duration, "ai_model": model,
        "outputs": {
            "single_mom_docx": docx_name,
            "complete_transcript": True,
            "complete_translation": True,
            "bilingual_document": True,
            "executive_summary": True,
            "complete_conversation_summary": True
        },
        "validation": report,
        "previous_meeting_context": final_ai.get(
            "previous_meeting_context", {"used_as_reference": False}),
        "warnings": warnings
    }), JSON_MIME)
    ctx.store.upsert_text(batches_folder, "BATCH_MANIFEST.json", dump({
        "meeting_id": env("MEETING_ID"), "duration_seconds": duration,
        "batch_seconds": BATCH_SECONDS, "batch_total": len(files),
        "status": "complete",
        "batches": [b.get("batch") for b in batch_results]
    }), JSON_MIME)
    for p in [ctx.p["source"], ctx.p["wav"]]:
        p.unlink(missing_ok=True)
    st.update("COMPLETED",
              f"Complete meeting processed in {len(files)} batches. Single MoM ready.",
              100, batch_index=len(files), batch_total=len(files))
